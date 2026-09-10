import os
from typing import Dict, Any, Optional
import fitz  # PyMuPDF
import docx
import openpyxl
from pptx import Presentation
from PIL import Image
import pytesseract
import config

pytesseract.pytesseract.tesseract_cmd = config.TESSERACT_CMD


def calculate_ocr_confidence(data_dict: dict) -> Optional[float]:
    confs = [int(c) for c in data_dict.get("conf", []) if int(c) >= 0]
    if confs:
        return round(sum(confs) / len(confs), 2)
    return None


def extract_text_from_file(file_path: str, filename: str) -> Dict[str, Any]:
    ext = os.path.splitext(filename)[1].lower()

    if ext in [".png", ".jpg", ".jpeg"]:
        try:
            img = Image.open(file_path)
            text = pytesseract.image_to_string(img)
            data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
            confidence = calculate_ocr_confidence(data)
            return {
                "extracted_text": text.strip(),
                "used_ocr": True,
                "ocr_confidence": confidence,
            }
        except Exception as e:
            raise ValueError(f"Image OCR processing failed: {str(e)}")

    elif ext == ".txt":
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
            return {
                "extracted_text": text.strip(),
                "used_ocr": False,
                "ocr_confidence": None,
            }
        except Exception as e:
            raise ValueError(f"TXT extraction failed: {str(e)}")

    elif ext == ".pdf":
        try:
            doc = fitz.open(file_path)
            extracted_pages = []
            used_ocr = False
            confidences = []

            for page in doc:
                text = page.get_text()
                if len(text.strip()) > 20:
                    extracted_pages.append(text.strip())
                else:
                    used_ocr = True
                    pix = page.get_pixmap(dpi=150)
                    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    ocr_text = pytesseract.image_to_string(img)
                    data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
                    conf = calculate_ocr_confidence(data)
                    if conf is not None:
                        confidences.append(conf)
                    extracted_pages.append(ocr_text.strip())

            full_text = "\n\n".join(extracted_pages)
            avg_conf = (
                round(sum(confidences) / len(confidences), 2) if confidences else None
            )

            return {
                "extracted_text": full_text.strip(),
                "used_ocr": used_ocr,
                "ocr_confidence": avg_conf,
            }
        except Exception as e:
            raise ValueError(f"PDF extraction failed: {str(e)}")

    elif ext == ".docx":
        try:
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n".join(paragraphs)
            return {
                "extracted_text": full_text.strip(),
                "used_ocr": False,
                "ocr_confidence": None,
            }
        except Exception as e:
            raise ValueError(f"DOCX extraction failed: {str(e)}")

    elif ext == ".xlsx":
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            lines = []
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                lines.append(f"--- Sheet: {sheetname} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(val) for val in row if val is not None]
                    if row_vals:
                        lines.append(" | ".join(row_vals))
            full_text = "\n".join(lines)
            return {
                "extracted_text": full_text.strip(),
                "used_ocr": False,
                "ocr_confidence": None,
            }
        except Exception as e:
            raise ValueError(f"XLSX extraction failed: {str(e)}")

    elif ext == ".pptx":
        try:
            prs = Presentation(file_path)
            lines = []
            for i, slide in enumerate(prs.slides, 1):
                lines.append(f"--- Slide {i} ---")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        lines.append(shape.text.strip())
            full_text = "\n".join(lines)
            return {
                "extracted_text": full_text.strip(),
                "used_ocr": False,
                "ocr_confidence": None,
            }
        except Exception as e:
            raise ValueError(f"PPTX extraction failed: {str(e)}")

    else:
        raise ValueError(
            f"Unsupported file format '{ext}'. Supported formats: PDF, DOCX, XLSX, PPTX, TXT, PNG, JPG, JPEG"
        )
